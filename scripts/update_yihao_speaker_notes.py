#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation


NOTES = {
    1: (
        "开场（20-30秒）\n"
        "大家好，我是陈奕好。这次汇报聚焦于医疗影像与AI在介入导航中的应用，"
        "我会重点展示我近期在冠脉分析pipeline上的工程化推进和量化结果。"
    ),
    2: (
        "议程（20秒）\n"
        "先简要介绍我的背景和能力，再进入当前核心项目 vessel_seg，"
        "最后给出风险、下一步计划，以及希望老师给我的反馈点。"
    ),
    3: (
        "个人能力（35-45秒）\n"
        "我的强项是把问题从算法点推进到可复现的系统流程：\n"
        "1) 能完整搭建 data -> preprocessing -> model -> evaluation。\n"
        "2) 能结合几何约束与学习方法处理医疗影像问题。\n"
        "3) 在工程上注重可复现和可审计，便于快速迭代。"
    ),
    4: (
        "经历地图（30-40秒）\n"
        "这页展示我的项目积累路径：从硬件/系统实现，到医学影像建模与评估。"
        "核心不是做很多零散demo，而是逐步形成面向真实研究问题的稳定推进能力。"
    ),
    5: (
        "实现经验1（30-40秒）\n"
        "在硬件与系统控制项目中，我负责采集与控制链路实现，"
        "重点是多传感器协同和可重复验证。这个阶段训练了我做复杂系统调试和故障定位的能力。"
    ),
    6: (
        "实现经验2（30-40秒）\n"
        "这里强调机械与结构层面的迭代思路：先考虑制造约束，再做几何细化和原型反馈闭环。"
        "这套方法后来直接迁移到我在医学影像pipeline上的工程迭代里。"
    ),
    7: (
        "过渡页（35-45秒）\n"
        "目前我已经把冠脉相关流程打通为可运行原型，并统一了输出结构。"
        "下一阶段会更强调两点：真实临床场景约束，以及分段标注/量化评估的一致性。"
    ),
    8: (
        "项目总览（40-50秒）\n"
        "这是 vessel_seg 的5步流程：分割、中心线提取、修复、特征提取、重建对比。"
        "我把每一步都定义了量化指标，目标是把问题从“看起来不错”变成“可测量、可比较、可回归”。"
    ),
    9: (
        "Step1 正例（40秒）\n"
        "Normal_1 的 Dice 和 HD95 都比较好，说明分割质量稳定。"
        "这个case证明当前上游质量足以支撑后续中心线和特征分析。"
    ),
    10: (
        "Step1 难例（45秒）\n"
        "Normal_2 的 HD95 偏高，提示存在明显边界误差。"
        "这说明 Step1 稳定性是全流程上限，若上游波动大，下游再优化也会受限。"
    ),
    11: (
        "Step3 对比（45-60秒）\n"
        "这里看中心线在修复前后的几何质量和对GT贴合度。"
        "关注点是：误差尾部是否下降、连通性是否更符合血管解剖结构。"
        "这一步是后续特征和重建可靠性的关键。"
    ),
    12: (
        "关键修复（60秒）\n"
        "我定位到一个拓扑层面的根因：端点在写VTP时重复建点，导致视觉像连着、拓扑却是碎的。"
        "修复后改为共享端点ID，并加入拓扑监控字段，"
        "把问题从临时patch变成可验证、可回归的工程修复。"
    ),
    13: (
        "跨case总结（50-60秒）\n"
        "这一页是关键决策页：横向比较不同case，判断瓶颈到底在Step1、Step3还是Step4。"
        "目前结论是：Step1稳定性和Step4分支治理是优先优化方向。"
    ),
    14: (
        "Step4（旧版页，可作为备份，30-40秒）\n"
        "这一页展示分支级特征提取与GT对比。若时间有限可以快讲，"
        "重点只说两句：分支数量一致性和描述子相似度。"
    ),
    15: (
        "风险与计划（旧版页，可作为备份，30-40秒）\n"
        "如果保留此页，强调三点风险：分支计数偏差、case波动、回归集不足；"
        "并说明你已有对应两周执行计划。"
    ),
    16: (
        "贡献总结（旧版页，可作为备份，25-35秒）\n"
        "若讲到此页，聚焦工程价值：端到端可复现、统一量化口径、关键bug闭环修复。"
    ),
    17: (
        "Step4（新版主讲页，50秒）\n"
        "这里重点讲“分支层面的可解释比较”："
        "branch count diff反映拓扑是否对齐，descriptor cosine反映形态特征是否接近GT。"
        "当前结果说明整体方向正确，但仍需压缩冗余分支。"
    ),
    18: (
        "风险与两周计划（60秒）\n"
        "风险：Step4偶发分支偏差、跨case波动、失败样本覆盖不足。\n"
        "计划：Week1 做 repair 触发与拓扑约束强化；Week2 做分支治理+10-20例批评估；"
        "输出可复现图表和对比报告，直接支撑导师讨论与后续投稿路线。"
    ),
    19: (
        "核心贡献（45-55秒）\n"
        "我这阶段最核心的价值不是单个模型指标，而是把研究问题工程化："
        "统一流程、统一指标、统一回归机制。"
        "这让每次改动都能被量化验证，能稳定推进而不是反复试错。"
    ),
    20: (
        "收尾提问（45秒）\n"
        "我希望老师给三个方向性反馈：\n"
        "1) 指标优先级是否以Step1稳定性为先；\n"
        "2) 论文叙事更偏方法创新还是工程闭环；\n"
        "3) 小样本闭环到批量验证的节奏是否合理。\n"
        "谢谢老师，欢迎直接指出优先级调整建议。"
    ),
}


def select_input(default_a: Path, default_b: Path) -> Path:
    if default_a.exists():
        return default_a
    if default_b.exists():
        return default_b
    raise FileNotFoundError(f"Neither {default_a} nor {default_b} exists")


def main() -> None:
    parser = argparse.ArgumentParser(description="Update Yihao PPT speaker notes for all slides.")
    parser.add_argument(
        "--infile",
        type=Path,
        default=None,
        help="Input PPT path. If omitted, try Applications/Yihao_last4_updated.pptx then Applications/Yihao.pptx",
    )
    parser.add_argument(
        "--outfile",
        type=Path,
        default=Path("Applications/Yihao_with_full_notes.pptx"),
        help="Output PPT path.",
    )
    args = parser.parse_args()

    infile = args.infile
    if infile is None:
        infile = select_input(Path("Applications/Yihao_last4_updated.pptx"), Path("Applications/Yihao.pptx"))
    if not infile.exists():
        raise FileNotFoundError(infile)

    prs = Presentation(str(infile))
    total = len(prs.slides)
    for idx, slide in enumerate(prs.slides, start=1):
        note = NOTES.get(idx, "")
        notes_frame = slide.notes_slide.notes_text_frame
        notes_frame.clear()
        notes_frame.text = note

    args.outfile.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(args.outfile))
    print(f"[done] wrote {args.outfile}")
    print(f"[info] source={infile} slides={total}")


if __name__ == "__main__":
    main()
