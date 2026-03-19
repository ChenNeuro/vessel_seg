#!/usr/bin/env python3
from __future__ import annotations

import argparse
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from typing import Iterable, List

from pptx import Presentation


def remove_slide(prs: Presentation, index: int) -> None:
    sld_id_lst = prs.slides._sldIdLst  # type: ignore[attr-defined]
    sld = sld_id_lst[index]
    rel_id = sld.rId
    prs.part.drop_rel(rel_id)
    del sld_id_lst[index]


def parse_slide_indices(spec: str) -> List[int]:
    result: List[int] = []
    for part in spec.split(","):
        p = part.strip()
        if not p:
            continue
        if "-" in p:
            start_s, end_s = p.split("-", 1)
            start = int(start_s)
            end = int(end_s)
            if end < start:
                start, end = end, start
            result.extend(range(start, end + 1))
        else:
            result.append(int(p))
    if not result:
        raise ValueError("empty source slide spec")
    uniq = sorted(set(result))
    return uniq


def copy_slide_content(src_slide, dst_slide) -> None:
    for shape in src_slide.shapes:
        if shape.shape_type == 13:  # PICTURE
            image_blob = shape.image.blob
            stream = BytesIO(image_blob)
            dst_slide.shapes.add_picture(stream, shape.left, shape.top, shape.width, shape.height)
        else:
            newel = deepcopy(shape.element)
            dst_slide.shapes._spTree.insert_element_before(newel, "p:extLst")


def delete_last_n_slides(prs: Presentation, n: int) -> None:
    if n <= 0:
        return
    if n > len(prs.slides):
        raise ValueError(f"cannot delete {n} slides from deck with {len(prs.slides)} slides")
    for idx in range(len(prs.slides) - 1, len(prs.slides) - n - 1, -1):
        remove_slide(prs, idx)


def ensure_file(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Replace the last N slides of a target deck with selected slides "
            "from a source deck."
        )
    )
    parser.add_argument(
        "--target",
        type=Path,
        default=Path("Applications/Yihao.pptx"),
        help="PPT to update.",
    )
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("docs/teacher_report_20260210_slides_latest.pptx"),
        help="Source PPT that contains fresh report slides.",
    )
    parser.add_argument(
        "--source-slides",
        type=str,
        default="9-12",
        help="1-based source slide indices, e.g. '9-12' or '9,10,11,12'.",
    )
    parser.add_argument(
        "--replace-count",
        type=int,
        default=4,
        help="How many slides to remove from the end of target.",
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=Path("Applications/Yihao_last4_updated.pptx"),
        help="Output file path.",
    )
    args = parser.parse_args()

    ensure_file(args.target)
    ensure_file(args.source)

    src_indices = parse_slide_indices(args.source_slides)
    target = Presentation(str(args.target))
    source = Presentation(str(args.source))

    for idx in src_indices:
        if idx < 1 or idx > len(source.slides):
            raise IndexError(f"source slide index out of range: {idx}")

    delete_last_n_slides(target, args.replace_count)
    layout = target.slide_layouts[0]
    for idx in src_indices:
        src_slide = source.slides[idx - 1]
        dst_slide = target.slides.add_slide(layout)
        copy_slide_content(src_slide, dst_slide)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    target.save(str(args.out))
    print(f"[done] wrote {args.out}")
    print(f"[info] target slides now: {len(target.slides)}")


if __name__ == "__main__":
    main()
