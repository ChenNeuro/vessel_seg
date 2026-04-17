"""PPT helpers for the repo-level modeling flow presentation."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

BG = RGBColor(248, 249, 244)
TEXT = RGBColor(26, 32, 44)
MUTED = RGBColor(92, 102, 114)
LINE = RGBColor(120, 133, 145)
ACCENT_RED = RGBColor(221, 89, 79)
ACCENT_BLUE = RGBColor(74, 109, 167)
ACCENT_GREEN = RGBColor(88, 144, 114)
ACCENT_GOLD = RGBColor(196, 151, 64)
PANEL = RGBColor(255, 255, 255)
PANEL_ALT = RGBColor(239, 244, 250)
PANEL_WARM = RGBColor(253, 245, 232)
PANEL_SOFT = RGBColor(237, 247, 240)


@dataclass(frozen=True)
class TeacherModelingPptConfig:
    """Configuration for the modeling-flow teacher presentation."""

    output_path: Path
    notes_path: Path | None = None
    title: str = "冠脉建模流程与所需数据"
    subtitle: str = "基于仓库现状，面向“现有分割 -> 世界系放置 -> X 光重建”的汇报版本"
    pipeline_overview_png: Path | None = None
    cross_case_summary_png: Path | None = None
    normal1_centerline_png: Path | None = None
    normal2_centerline_png: Path | None = None


def _set_slide_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def _add_textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    font_size: int,
    bold: bool = False,
    color: RGBColor = TEXT,
    align=PP_ALIGN.LEFT,
    margin_left: int = 8,
    margin_right: int = 8,
    margin_top: int = 4,
    margin_bottom: int = 4,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Pt(margin_left)
    frame.margin_right = Pt(margin_right)
    frame.margin_top = Pt(margin_top)
    frame.margin_bottom = Pt(margin_bottom)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bool(bold)
    run.font.color.rgb = color


def _add_title(slide, title: str, subtitle: str) -> None:
    _set_slide_background(slide)
    _add_textbox(slide, 0.65, 0.38, 11.9, 0.58, title, font_size=24, bold=True)
    _add_textbox(slide, 0.67, 0.95, 11.7, 0.38, subtitle, font_size=11, color=MUTED)
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.65), Inches(1.28), Inches(12.0), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT_RED
    line.line.fill.background()


def _add_panel(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    fill: RGBColor = PANEL,
    line: RGBColor = LINE,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.0)


def _add_box_text(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
    body: str,
    *,
    fill: RGBColor,
    accent: RGBColor,
) -> None:
    _add_panel(slide, left, top, width, height, fill=fill, line=accent)
    _add_textbox(slide, left + 0.12, top + 0.06, width - 0.24, 0.34, title, font_size=14, bold=True, color=accent)
    _add_textbox(slide, left + 0.12, top + 0.42, width - 0.24, height - 0.5, body, font_size=11, color=TEXT)


def _add_bullets(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    title: str,
    bullets: list[str],
    *,
    fill: RGBColor = PANEL,
    accent: RGBColor = ACCENT_BLUE,
) -> None:
    _add_panel(slide, left, top, width, height, fill=fill, line=accent)
    _add_textbox(slide, left + 0.12, top + 0.06, width - 0.24, 0.3, title, font_size=14, bold=True, color=accent)
    box = slide.shapes.add_textbox(Inches(left + 0.12), Inches(top + 0.38), Inches(width - 0.24), Inches(height - 0.46))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Pt(6)
    frame.margin_right = Pt(4)
    frame.margin_top = Pt(2)
    frame.margin_bottom = Pt(2)
    for idx, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if idx == 0 else frame.add_paragraph()
        paragraph.text = f"• {bullet}"
        paragraph.level = 0
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.font.size = Pt(11)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(3)


def _add_arrow(slide, x1: float, y1: float, x2: float, y2: float, *, color: RGBColor = LINE, width: float = 1.8) -> None:
    connector = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(width)
    try:
        connector.line.end_arrowhead = True
    except Exception:
        pass


def _add_footer(slide, text: str) -> None:
    _add_textbox(slide, 0.75, 7.0, 11.8, 0.22, text, font_size=8, color=MUTED)


def _add_table(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(left), Inches(top), Inches(width), Inches(height)).table
    column_widths = [width * 0.27, width * 0.20, width * 0.53]
    for idx, column_width in enumerate(column_widths):
        if idx < len(headers):
            table.columns[idx].width = Inches(column_width)

    header_fill = ACCENT_BLUE
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = header_fill
        paragraph = cell.text_frame.paragraphs[0]
        paragraph.font.size = Pt(11)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(255, 255, 255)
        paragraph.alignment = PP_ALIGN.CENTER

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, value in enumerate(row):
            cell = table.cell(row_idx, col_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if row_idx % 2 == 1 else PANEL_ALT
            paragraph = cell.text_frame.paragraphs[0]
            paragraph.font.size = Pt(10)
            paragraph.font.color.rgb = TEXT
            paragraph.alignment = PP_ALIGN.LEFT


def _image_exists(path: Path | None) -> bool:
    return bool(path is not None and path.exists())


def _build_title_slide(slide, config: TeacherModelingPptConfig) -> None:
    _add_title(slide, config.title, config.subtitle)
    _add_bullets(
        slide,
        0.85,
        1.75,
        4.25,
        2.15,
        "汇报目标",
        [
            "把当前仓库已经实现的工程主线，与老师提出的“术中重建”思路对齐。",
            "回答三个问题：现在做到哪里、下一步建模怎么写、需要补什么数据。",
            "核心口径：先消费现有分割/clean tree，再放到世界系里，最后输出可比较的 X 光重建。",
        ],
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_bullets(
        slide,
        5.35,
        1.75,
        3.35,
        2.15,
        "仓库里已经有的能力",
        [
            "五阶段工程主线",
            "语义拓扑与左右系统划分",
            "分支聚类与 prototype prior",
            "FGPM / Bayesian 形状建模接口",
            "C-arm 几何与 synthetic projection",
        ],
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_bullets(
        slide,
        8.95,
        1.75,
        3.35,
        2.15,
        "老师这版重点",
        [
            "参考系设定",
            "状态节点与观测量",
            "聚类能否变成抽象结构体",
            "Bayesian 估计能否接上真实 X 光",
            "构型空间和工作空间如何区分",
        ],
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_box_text(
        slide,
        0.9,
        4.35,
        11.35,
        1.75,
        "一句话总览",
        "当前仓库已经具备“CT / mask -> centerline tree -> semantic topology -> branch features -> render”的工程闭环；老师想推进的，是把这条静态结构链升级成“世界系 + 心脏系 + C 臂观测系”下的参数化状态空间模型，再把 synthetic X-ray 与真实 X 光做匹配。",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_footer(slide, "依据仓库：README、docs/coronary_analysis_master_blueprint.md、vessel_seg/pipeline/*、vessel_seg/reconstruction_3d2d/*")


def _build_repo_review_slide(slide) -> None:
    _add_title(slide, "Repo Review：当前代码已经分成五条主线", "这一页用来证明 PPT 不是空想，而是直接从仓库现状抽出来的")
    _add_box_text(
        slide,
        0.8,
        1.65,
        2.3,
        1.55,
        "1. 工程主线",
        "vessel_seg/pipeline/\n五阶段统一入口，负责每例 case 的标准输出。",
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_box_text(
        slide,
        3.25,
        1.65,
        2.3,
        1.55,
        "2. 语义拓扑",
        "semantic_topology.py\n病例内解剖坐标系、左右系统、主干家族与语义赋值。",
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_box_text(
        slide,
        5.7,
        1.65,
        2.3,
        1.55,
        "3. 聚类原型",
        "branch_clustering.py\nside-first clustering，把 branch 组织成 prototype graph。",
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_box_text(
        slide,
        8.15,
        1.65,
        2.3,
        1.55,
        "4. 概率形状",
        "fgpm.py + tree_prior.py\nFourier + GP + MAP，为 Bayesian 形状先验留了接口。",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_box_text(
        slide,
        10.6,
        1.65,
        1.95,
        1.55,
        "5. 3D-2D",
        "reconstruction_3d2d/\nHeartState、CArmConfig、synthetic projection。",
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_bullets(
        slide,
        0.95,
        3.55,
        5.4,
        2.5,
        "对应老师的讨论点",
        [
            "“参考系设定”已经有雏形：HeartState + CArmConfig + W0 / B / C / O / H 变换链。",
            "“聚类估计”已经有 side-first baseline，但现在输出仍主要是分析报告，下一步应升级成 prototype prior。",
            "“贝叶斯估计”已有 FGPM / MAP 形状建模接口，但尚未与 X 光匹配串成在线估计。",
        ],
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_bullets(
        slide,
        6.65,
        3.55,
        5.6,
        2.5,
        "这意味着什么",
        [
            "仓库不是从零开始做重建，而是已经有“静态结构模型”和“简化观测模型”两头。",
            "真正缺的是中间那层：状态参数化、类手术数据构造、真实 X 光 matching、时序 Bayesian 更新。",
            "因此本次 PPT 应该强调“从现有工程主线自然生长到术中重建层”，而不是另起炉灶。",
        ],
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_footer(slide, "关键文件：vessel_seg/pipeline/stages.py、vessel_seg/pipeline/semantic_topology.py、vessel_seg/pipeline/branch_clustering.py、vessel_seg/fgpm.py、vessel_seg/reconstruction_3d2d/contracts.py")


def _build_pipeline_slide(slide, config: TeacherModelingPptConfig) -> None:
    _add_title(slide, "当前已经可复现的工程主线", "先把已有 pipeline 讲清楚，再接老师希望看到的建模层")
    if _image_exists(config.pipeline_overview_png):
        slide.shapes.add_picture(str(config.pipeline_overview_png), Inches(0.8), Inches(1.55), width=Inches(7.2))
    else:
        _add_box_text(
            slide,
            0.8,
            1.55,
            7.1,
            3.95,
            "Pipeline 概览图缺失",
            "默认会读取 docs/archive/assets/teacher_report_20260209/pipeline_overview.png。",
            fill=PANEL_WARM,
            accent=ACCENT_GOLD,
        )
    _add_bullets(
        slide,
        8.15,
        1.55,
        4.25,
        4.0,
        "每个 stage 的关键产物",
        [
            "Step1 CT segmentation：mask.nii.gz",
            "Step2 centerline extraction：centerlines.vtp + branch manifest",
            "Step3 centerline repair：tree.json + branch_names.json + semantic_topology.json",
            "Step4 wall features：branch_dataset.npz / radii profiles",
            "Step5 rendering：overview.png / case summary",
        ],
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_box_text(
        slide,
        0.95,
        5.85,
        11.3,
        0.85,
        "这一页的结论",
        "老师后面要的术中重建，不需要推翻现有主线；它应该直接消费 Step3/Step4 的 clean tree、语义拓扑和半径信息，把这些静态产物重写成状态空间里的先验与观测模型。",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_footer(slide, "代码入口：vessel_seg/pipeline/orchestrator.py、vessel_seg/pipeline/stages.py；CLI：python -m vessel_seg pipeline-case")


def _build_flow_mapping_slide(slide) -> None:
    _add_title(slide, "老师想法与当前仓库的映射", "把“现有分割 -> 世界系放置 -> 输出 X 光重建信息”写成一条清晰流程")
    box_specs = [
        (0.7, 1.95, 2.0, 0.95, "输入层", "CT / mask\n已有 X 光 / ECG"),
        (3.0, 1.95, 2.1, 0.95, "结构层", "clean tree\nsemantic topology"),
        (5.45, 1.95, 2.1, 0.95, "原型层", "cluster ->\nprototype prior"),
        (7.9, 1.95, 2.1, 0.95, "状态层", "[x,theta]_cav,L,R\nbeta_soft, ecg"),
        (10.35, 1.95, 2.1, 0.95, "观测层", "world pose + C-arm\n3D -> 2D projection"),
    ]
    fills = [PANEL, PANEL_ALT, PANEL_WARM, PANEL_SOFT, PANEL_ALT]
    accents = [ACCENT_RED, ACCENT_BLUE, ACCENT_GOLD, ACCENT_GREEN, ACCENT_BLUE]
    for (left, top, width, height, title, body), fill, accent in zip(box_specs, fills, accents):
        _add_box_text(slide, left, top, width, height, title, body, fill=fill, accent=accent)
    _add_arrow(slide, 2.72, 2.42, 2.98, 2.42, color=ACCENT_RED)
    _add_arrow(slide, 5.12, 2.42, 5.43, 2.42, color=ACCENT_BLUE)
    _add_arrow(slide, 7.57, 2.42, 7.88, 2.42, color=ACCENT_GOLD)
    _add_arrow(slide, 10.02, 2.42, 10.33, 2.42, color=ACCENT_GREEN)
    _add_box_text(
        slide,
        3.5,
        3.45,
        2.0,
        0.88,
        "实验层",
        "采样 config / space\n合成类手术数据",
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_box_text(
        slide,
        8.3,
        3.45,
        2.0,
        0.88,
        "实时层",
        "真实 X 光 vs 模型投影\nmatching / MAP",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_arrow(slide, 4.48, 3.43, 4.48, 2.93, color=ACCENT_GOLD, width=1.5)
    _add_arrow(slide, 9.28, 3.43, 9.28, 2.93, color=ACCENT_GREEN, width=1.5)
    _add_bullets(
        slide,
        0.85,
        4.85,
        11.4,
        1.55,
        "要点",
        [
            "现在仓库已经覆盖“结构层 + 部分观测层”，老师要推进的是“原型层 + 状态层 + 实时匹配层”。",
            "最小闭环应先做 synthetic X-ray，再把真实 X 光作为观测输入，逐步过渡到 Bayesian / MAP 估计。",
        ],
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_footer(slide, "与老师表述对齐：先利用现有分割，在世界系放置，最后输出 X 光重建信息，此时参数可变")


def _build_reference_slide(slide) -> None:
    _add_title(slide, "参考系设定：W0 / B / H / C / O", "这一页回答“手术台参数、C 臂角度和相对位姿、投影变换”")
    _add_box_text(
        slide,
        0.8,
        1.7,
        3.4,
        3.0,
        "床零位世界系 W0 + 当前床台系 B",
        "W0 负责给出床零位统一基准；B 负责显式表示床体 longitudinal / lateral / height / tilt。\n\n老师说的“手术台参数”，建议统一写进 ^W0T_B 这一层。",
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_box_text(
        slide,
        4.55,
        1.7,
        3.1,
        3.0,
        "心脏系 H",
        "推荐由 apex / base center / LM ostium / LM bifurcation 定义解剖坐标系。\n\n状态量：coronary frame residual、[x,theta]_cav、[x,theta]_L、[x,theta]_R、beta_soft。\n\n这是把静态树升级成参数化结构的核心。",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_box_text(
        slide,
        8.0,
        1.7,
        4.0,
        3.0,
        "机械系 C + 观测系 O",
        "C 由 alpha(LAO/RAO)、beta(CRA/CAU) 给出机械姿态；O 由 source_to_isocenter、source_to_detector、pixel spacing、principal point、distortion 定义。\n\n仓库里已经统一成 CArmConfig / DetectorConfig / observer model。",
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_arrow(slide, 4.18, 3.18, 4.52, 3.18, color=ACCENT_RED)
    _add_arrow(slide, 7.67, 3.18, 7.98, 3.18, color=ACCENT_BLUE)
    _add_box_text(
        slide,
        1.0,
        5.2,
        10.9,
        0.95,
        "建议的数据约定",
        "所有 tree / branch 几何先在 H 表达；投影时再走 ^B T_H、^W0 T_B、^C T_W0、^O T_C。这样可以把“构型变化”“床台变化”和“成像几何变化”拆开处理。",
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_footer(slide, "对应代码：vessel_seg/reconstruction_3d2d/contracts.py、vessel_seg/reconstruction_3d2d/carm_geometry.py")


def _build_data_tree_slide(slide) -> None:
    _add_title(slide, "数据树：构型空间与工作空间怎么分", "这一页直接回答“数据树问题、构型空间和工作空间的问题”")
    _add_panel(slide, 0.8, 1.7, 5.15, 4.95, fill=PANEL, line=ACCENT_RED)
    _add_textbox(slide, 1.0, 1.9, 4.75, 0.35, "建议的数据树", font_size=15, bold=True, color=ACCENT_RED)
    tree_text = (
        "Input\n"
        "├─ ECG_t\n"
        "├─ Xray_t\n"
        "└─ Preop CT / mask / clean tree\n\n"
        "World W0 / Bed B\n"
        "├─ ^W0T_B(longitudinal, lateral, height, tilt)\n"
        "└─ ^B T_H\n\n"
        "C-arm / Observer\n"
        "├─ CArmConfig(alpha, beta, d_SD, d_SI, detector)\n"
        "└─ ^C T_W0, ^O T_C\n\n"
        "Heart H\n"
        "├─ [x,theta]_cav\n"
        "├─ [x,theta]_L\n"
        "├─ [x,theta]_R\n"
        "└─ beta_soft / latent shape coeffs"
    )
    _add_textbox(slide, 1.05, 2.28, 4.65, 4.0, tree_text, font_size=11)
    _add_bullets(
        slide,
        6.25,
        1.7,
        2.85,
        2.25,
        "构型空间 config space",
        [
            "低维参数：^W0T_B、^B T_H、q_L、q_R、beta_soft、ecg_phase、c_arm。",
            "适合采样、优化、Bayesian filtering。",
            "老师说的 [config, space] 可以先从这里落地。",
        ],
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_bullets(
        slide,
        9.35,
        1.7,
        2.85,
        2.25,
        "工作空间 workspace",
        [
            "3D 中心线 / 分支半径 / mesh / 世界坐标中的几何。",
            "2D detector 上的投影点、投影宽度、mask。",
            "这里是合成观测和真实观测发生比较的地方。",
        ],
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_bullets(
        slide,
        6.25,
        4.2,
        5.95,
        2.45,
        "为什么要分开",
        [
            "如果不区分 config 和 workspace，手术台位姿、心脏运动、血管形变会全部混在一个坐标里，难以解释也难以估计。",
            "分开以后，聚类和先验只约束 config space；投影、成像误差和遮挡则主要发生在 workspace。",
        ],
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_footer(slide, "建议在 JSON / dataclass 里显式区分 state params 与 geometric observations")


def _build_clustering_slide(slide) -> None:
    _add_title(slide, "聚类估计：能不能做成抽象结构体？", "这一页回答“换数据结构、隐形结构、软性结构，能不能够”")
    _add_box_text(
        slide,
        0.8,
        1.75,
        3.65,
        1.55,
        "显式结构层",
        "tree.json\n{parent, lambda, theta, phi, depth, child_count}\n\n这是当前仓库最稳定的结构基础。",
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_box_text(
        slide,
        4.85,
        1.75,
        3.65,
        1.55,
        "抽象原型层",
        "side-first clustering + canonical naming\n输出 prototype graph：trunk / continuation / side / terminal。",
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_box_text(
        slide,
        8.9,
        1.75,
        3.45,
        1.55,
        "软性隐变量层",
        "FGPM / soft_coeffs / deformation latent\n表达不可直接观测但会影响投影的柔性形变。",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_arrow(slide, 4.47, 2.52, 4.82, 2.52, color=ACCENT_RED)
    _add_arrow(slide, 8.52, 2.52, 8.87, 2.52, color=ACCENT_GOLD)
    _add_bullets(
        slide,
        0.95,
        3.75,
        5.45,
        2.35,
        "当前仓库已经支持的部分",
        [
            "semantic_topology.py：给分支加上语义家族和病例内解剖坐标。",
            "branch_clustering.py：先分 LCA / RCA，再在各自内部聚类，避免左右混簇。",
            "fgpm.py：Fourier + GP + MAP，已经说明软性形状先验是可以建的。",
        ],
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_bullets(
        slide,
        6.65,
        3.75,
        5.55,
        2.35,
        "建议在 PPT 里的表述",
        [
            "聚类不应该只停留在“看 gallery”，而要输出一个 prototype prior，作为状态空间的结构约束。",
            "可以先把每个 cluster 看成一个抽象结构体槽位，再在槽位里挂上长度、方向、半径和 soft latent。",
            "这样就能回答老师问的“能不能实现类手术的数据构造”：可以，先在原型图上采样，再映射到具体 geometry。",
        ],
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_footer(slide, "依据文档：docs/branch_clustering_side_first_contract.md、docs/branch_clustering_lr_baseline.md")


def _build_bayesian_slide(slide) -> None:
    _add_title(slide, "状态节点、观测量与贝叶斯估计", "这一页把“模型 - 演化/仿真 - 实时匹配”写成数学语言")
    _add_bullets(
        slide,
        0.8,
        1.7,
        3.75,
        2.45,
        "状态 z_t",
        [
            "^W0T_B 与 ^B T_H：床台与心脏在 W0/B/H 链中的位姿",
            "q_cav / q_L / q_R：主干与左右冠的构型参数",
            "beta_soft：柔性形变低维系数",
            "phi_ecg：心动周期相位",
        ],
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_bullets(
        slide,
        4.8,
        1.7,
        3.65,
        2.45,
        "观测 y_t",
        [
            "真实 X 光图像或其 vessel segmentation",
            "ECG 相位 / 时间戳",
            "C 臂角度与 detector 标定",
            "可选：人工标注的关键点或血管对应关系",
        ],
        fill=PANEL_ALT,
        accent=ACCENT_BLUE,
    )
    _add_bullets(
        slide,
        8.7,
        1.7,
        3.55,
        2.45,
        "估计目标",
        [
            "forward：z_t -> synthetic X-ray",
            "inverse：真实 X 光 -> z_t 的后验",
            "先做 MAP matching，再扩展到 Bayesian filtering",
        ],
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_box_text(
        slide,
        1.05,
        4.45,
        11.0,
        1.15,
        "建议在汇报里直接给出的公式",
        "p(z_t | y_1:t) ∝ p(y_t | z_t) · p(z_t | z_{t-1}) · p(z_0)\n其中 p(y_t | z_t) 由 C 臂投影与图像 matching 给出，p(z_t | z_{t-1}) 由心动周期与柔性先验给出，p(z_0) 来自 prototype prior + FGPM。",
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_box_text(
        slide,
        1.05,
        5.9,
        11.0,
        0.78,
        "与仓库的关系",
        "HeartState / CArmConfig / synthetic_projection 已经提供了前向观测模型雏形；FGPM 提供了先验与 MAP 的语言；缺的是真实 X 光 matching 和时序更新。",
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_footer(slide, "对应老师关键词：状态节点、观测量、贝叶斯估计、模型演化/仿真")


def _build_data_requirements_slide(slide) -> None:
    _add_title(slide, "需要的数据与参数表", "这一页可以直接回答老师：为了把模型跑起来，我们具体还缺什么")
    headers = ["数据 / 参数", "仓库现状", "主要用途"]
    rows = [
        ["CCTA / coronary mask", "已有", "术前解剖入口；可生成 clean tree、wall features、shape prior。"],
        ["centerline tree + semantic topology", "已有", "提供 parent / lambda / theta / phi、左右系统、语义主干。"],
        ["branch radii / wall features", "已有", "提供血管宽度、纵向半径、FGPM 建模基础。"],
        ["C-arm 几何与 detector 标定", "接口已留，真实数据缺", "定义投影矩阵、pixel spacing、principal point、畸变。"],
        ["手术台 / 世界系位姿", "缺", "把心脏系映射到世界系，回答“患者怎么摆、器械怎么摆”。"],
        ["ECG 相位 / 时间戳", "缺", "建模心脏时序，连接 z_t 和 z_{t-1}。"],
        ["真实 X 光 / fluoroscopy", "缺", "作为观测 y_t，用于 matching / MAP / Bayes filtering。"],
        ["评估标注", "待定义", "关键点、分支对应、2D mask 或 2D-3D 注册基准。"],
    ]
    _add_table(slide, 0.8, 1.75, 11.75, 4.4, headers, rows)
    _add_box_text(
        slide,
        0.95,
        6.35,
        11.3,
        0.72,
        "优先补什么",
        "第一优先级不是再换分割模型，而是补齐真实 C 臂几何、ECG 相位和同步 X 光；只有这三类数据到位，真实 matching 和 Bayesian 估计才有入口。",
        fill=PANEL_WARM,
        accent=ACCENT_GOLD,
    )
    _add_footer(slide, "数据入口参考：data/README.md、data/external/asoca2020.path；新增术中数据建议单独建 manifest/schema")


def _build_evidence_and_roadmap_slide(slide, config: TeacherModelingPptConfig) -> None:
    _add_title(slide, "当前证据与下一步路线", "把“已经做到了什么”和“下一阶段做什么”收在一页里")
    if _image_exists(config.cross_case_summary_png):
        slide.shapes.add_picture(str(config.cross_case_summary_png), Inches(0.85), Inches(1.7), width=Inches(5.45))
    else:
        _add_box_text(
            slide,
            0.85,
            1.7,
            5.45,
            2.25,
            "量化总览图缺失",
            "默认读取 docs/archive/assets/teacher_report_20260209/cross_case_summary.png。",
            fill=PANEL_WARM,
            accent=ACCENT_GOLD,
        )
    if _image_exists(config.normal1_centerline_png):
        slide.shapes.add_picture(str(config.normal1_centerline_png), Inches(0.95), Inches(4.35), width=Inches(2.45))
    if _image_exists(config.normal2_centerline_png):
        slide.shapes.add_picture(str(config.normal2_centerline_png), Inches(3.55), Inches(4.35), width=Inches(2.45))
    _add_bullets(
        slide,
        6.65,
        1.7,
        5.4,
        2.15,
        "现在已经有的证据",
        [
            "Normal_1 / Normal_2 的五阶段 pipeline 可复现，并有量化指标和可视化资产。",
            "语义拓扑、左右系统分离、分支聚类 baseline 已经进入标准脚本与测试。",
            "synthetic projection 已经在代码层具备最小闭环接口。",
        ],
        fill=PANEL,
        accent=ACCENT_RED,
    )
    _add_bullets(
        slide,
        6.65,
        4.1,
        5.4,
        2.3,
        "建议的下一步",
        [
            "模型：把 cluster / prototype prior 和 FGPM 组合成可采样的类手术数据生成器。",
            "实验：系统采样 [config, space]，生成 [x,theta]_cav,L,R 的稀疏结构化表达与 synthetic X-ray 数据库。",
            "实时：定义真实 X 光与 synthetic X-ray 的 matching loss，再从单帧 MAP 过渡到时序 Bayesian filtering。",
        ],
        fill=PANEL_SOFT,
        accent=ACCENT_GREEN,
    )
    _add_footer(slide, "推荐汇报收束句：我们不是重新做一个分割项目，而是在现有冠脉结构化平台上补齐术中状态估计层")


def build_teacher_modeling_ppt(config: TeacherModelingPptConfig) -> Path:
    """Build the teacher-facing modeling-flow PPT."""
    config.output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT

    builders = [
        lambda slide: _build_title_slide(slide, config),
        _build_repo_review_slide,
        lambda slide: _build_pipeline_slide(slide, config),
        _build_flow_mapping_slide,
        _build_reference_slide,
        _build_data_tree_slide,
        _build_clustering_slide,
        _build_bayesian_slide,
        _build_data_requirements_slide,
        lambda slide: _build_evidence_and_roadmap_slide(slide, config),
    ]
    for builder in builders:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        builder(slide)

    presentation.save(str(config.output_path))
    return config.output_path


def build_teacher_modeling_notes(config: TeacherModelingPptConfig) -> Path | None:
    """Write concise speaker notes for the generated PPT."""
    if config.notes_path is None:
        return None
    config.notes_path.parent.mkdir(parents=True, exist_ok=True)
    content = """# 冠脉建模流程与所需数据：讲稿提纲

## Slide 1 标题页

- 先定口径：我们不是从零开始做术中重建，而是在现有冠脉结构化平台上继续往前走。
- 这次汇报要回答三件事：现有仓库做到哪里、老师要的建模层怎么接、真正还缺什么数据。

## Slide 2 Repo Review

- 仓库已经不是单一分割脚本，而是五条线并行：工程主线、语义拓扑、聚类原型、概率形状、3D-2D 观测。
- 这说明老师提出的参考系、聚类、Bayesian 估计并不是空中楼阁，代码里已经有接口。

## Slide 3 当前工程主线

- 先讲清楚五阶段 pipeline 是“静态结构层”。
- 强调 Step3 / Step4 最重要，因为术中重建会直接消费 clean tree、semantic topology 和 radii。

## Slide 4 老师想法映射

- 按输入层、结构层、原型层、状态层、观测层去讲。
- 明确指出：当前仓库已经有结构层和部分观测层，中间的状态建模和实时 matching 还要补。

## Slide 5 参考系

- 建议统一为床零位世界系 W0、当前床台系 B、心脏系 H、机械系 C、观测系 O。
- 手术台参数放在 ^W0T_B，心脏/冠脉刚体放在 ^B T_H，C 臂和 observer 几何分别放在 C/O 层。

## Slide 6 数据树 / Config vs Workspace

- config space 是低维参数，适合采样和估计。
- workspace 是真实 3D/2D 几何，适合投影和比较。
- 这样能把“结构变化”和“摆位变化”拆开。

## Slide 7 聚类与抽象结构体

- 回答老师：可以换成更抽象的数据结构。
- tree.json 是显式结构，prototype graph 是抽象结构，FGPM / soft coeffs 是软性隐变量。
- 接下来聚类的任务不应只是出图，而要输出可采样的原型先验。

## Slide 8 状态与贝叶斯估计

- 把问题写成 z_t 和 y_t。
- forward 是 synthetic X-ray，inverse 是真实 X 光匹配。
- 可以先讲 MAP matching，再说下一步推广到 Bayesian filtering。

## Slide 9 数据需求

- 术前 CT / mask / tree 这些已有。
- 真正缺的是术中同步数据：C 臂标定、世界系位姿、ECG、真实 X 光。
- 如果老师问“现在最大的缺口是什么”，答案就是术中数据和标定，而不是分割脚本。

## Slide 10 当前证据与路线

- 用现有量化结果证明平台不是概念验证。
- 再落到三步路线：模型采样、实验闭环、实时匹配。
- 最后一句可以收成：先做最小闭环 synthetic -> matching，再走到 online Bayesian 估计。

## 仓库依据

- `README.md`
- `docs/coronary_analysis_master_blueprint.md`
- `vessel_seg/pipeline/stages.py`
- `vessel_seg/pipeline/semantic_topology.py`
- `vessel_seg/pipeline/branch_clustering.py`
- `vessel_seg/fgpm.py`
- `vessel_seg/reconstruction_3d2d/contracts.py`
- `vessel_seg/reconstruction_3d2d/carm_geometry.py`
- `vessel_seg/reconstruction_3d2d/synthetic_projection.py`
"""
    config.notes_path.write_text(content, encoding="utf-8")
    return config.notes_path
