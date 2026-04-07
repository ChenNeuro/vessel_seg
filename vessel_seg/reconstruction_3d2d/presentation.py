"""PPT report helpers for 3D-2D reconstruction flow discussions."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


@dataclass(frozen=True)
class ReconstructionFlowPptConfig:
    """PPT generation contract for the 3D-2D reconstruction flow."""

    case_id: str
    output_path: Path
    title: str = "冠脉 3D-2D 重建流程图"
    subtitle: str = "从 clean tree 到 synthetic X-ray 的最小闭环"
    overview_png: Path | None = None
    projection_png: Path | None = None
    notes: tuple[str, ...] = ()


def _set_text(frame, text: str, *, font_size: int, bold: bool = False, color: RGBColor | None = None, align=PP_ALIGN.LEFT) -> None:
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bool(bold)
    if color is not None:
        run.font.color.rgb = color


def _add_title(slide, title: str, subtitle: str) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(11.8), Inches(0.8))
    _set_text(title_box.text_frame, title, font_size=24, bold=True, color=RGBColor(28, 28, 28))
    sub_box = slide.shapes.add_textbox(Inches(0.65), Inches(1.2), Inches(11.4), Inches(0.45))
    _set_text(sub_box.text_frame, subtitle, font_size=11, color=RGBColor(90, 90, 90))


def _add_box(slide, left: float, top: float, width: float, height: float, text: str, *, fill: RGBColor, line: RGBColor, font_size: int = 12) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1.25)
    _set_text(shape.text_frame, text, font_size=font_size, bold=True, color=RGBColor(20, 20, 20), align=PP_ALIGN.CENTER)


def _add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor) -> None:
    connector = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    connector.line.color.rgb = color
    connector.line.width = Pt(2.0)
    try:
        connector.line.end_arrowhead = True
    except Exception:
        pass


def _add_bullet_list(slide, left: float, top: float, width: float, height: float, title: str, bullets: list[str]) -> None:
    title_box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(0.35))
    _set_text(title_box.text_frame, title, font_size=16, bold=True, color=RGBColor(35, 35, 35))
    box = slide.shapes.add_textbox(Inches(left), Inches(top + 0.4), Inches(width), Inches(height - 0.4))
    frame = box.text_frame
    frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(12)
        paragraph.font.color.rgb = RGBColor(40, 40, 40)


def _build_flow_slide(slide, config: ReconstructionFlowPptConfig) -> None:
    _add_title(slide, "总流程图", "从现有 clean tree 出发，形成可估计的术中 3D-2D 重建链路")
    gray_fill = RGBColor(238, 242, 247)
    red_fill = RGBColor(255, 233, 233)
    blue_fill = RGBColor(225, 240, 255)
    green_fill = RGBColor(228, 246, 234)
    line = RGBColor(110, 125, 140)

    _add_box(slide, 0.6, 1.8, 2.0, 0.8, "输入层\nCCTA / 分割 / clean tree", fill=gray_fill, line=line)
    _add_box(slide, 3.0, 1.8, 2.0, 0.8, "结构层\nsemantic topology\n左右分离", fill=red_fill, line=line)
    _add_box(slide, 5.4, 1.8, 2.0, 0.8, "状态层\nHeartState\nq_L / q_R / beta", fill=blue_fill, line=line)
    _add_box(slide, 7.8, 1.8, 2.0, 0.8, "观测模型\nC-arm geometry\n3D -> 2D", fill=green_fill, line=line)
    _add_box(slide, 10.2, 1.8, 2.0, 0.8, "输出层\nsynthetic X-ray\nmatching", fill=gray_fill, line=line)

    _add_arrow(slide, 2.6, 2.2, 3.0, 2.2, RGBColor(86, 110, 135))
    _add_arrow(slide, 5.0, 2.2, 5.4, 2.2, RGBColor(86, 110, 135))
    _add_arrow(slide, 7.4, 2.2, 7.8, 2.2, RGBColor(86, 110, 135))
    _add_arrow(slide, 9.8, 2.2, 10.2, 2.2, RGBColor(86, 110, 135))

    _add_box(slide, 3.0, 3.4, 2.1, 0.75, "原型层\nPrototype graph\n聚类/先验", fill=RGBColor(255, 244, 214), line=line)
    _add_arrow(slide, 4.05, 3.35, 4.05, 2.6, RGBColor(160, 130, 55))

    _add_box(slide, 5.7, 3.4, 1.9, 0.75, "参数空间\nconfig space", fill=RGBColor(235, 230, 255), line=line)
    _add_box(slide, 8.0, 3.4, 1.9, 0.75, "工作空间\nworld / detector", fill=RGBColor(235, 250, 245), line=line)
    _add_arrow(slide, 7.6, 3.75, 8.0, 3.75, RGBColor(120, 120, 120))

    _add_bullet_list(
        slide,
        0.75,
        4.55,
        11.4,
        2.2,
        "这一页想表达的重点",
        [
            "五阶段主线提供 clean tree；3D-2D 重建层不替代主线，而是消费主线产物。",
            "聚类不再只是看图，而是输出 prototype prior，约束状态空间。",
            "术中问题被写成：状态 z_t 经投影模型 h(z_t) 生成观测，再与真实 X 光比较。",
        ],
    )


def _build_reference_slide(slide) -> None:
    _add_title(slide, "参考系、状态量与观测量", "把树结构升级成状态空间模型")
    _add_bullet_list(
        slide,
        0.7,
        1.6,
        3.8,
        4.9,
        "参考系",
        [
            "世界系 W：手术台、患者、C 臂统一参考系。",
            "心脏系 H：主动脉根/双开口附近的局部参考系。",
            "观察系 C：由 LAO/RAO、CRA/CAU、SID/SOD 决定的投影系。",
        ],
    )
    _add_bullet_list(
        slide,
        4.6,
        1.6,
        3.7,
        4.9,
        "状态量 z_t",
        [
            "T_WH：心脏在世界系的位姿。",
            "q_L / q_R：左右冠的运动学或段级参数。",
            "beta_soft：柔性形变低维系数。",
            "phi_ecg：心动周期相位。",
        ],
    )
    _add_bullet_list(
        slide,
        8.5,
        1.6,
        3.5,
        4.9,
        "观测量 y_t",
        [
            "X-ray 图像或其 vessel segmentation。",
            "ECG 相位信息。",
            "C 臂角度和几何参数。",
            "目标是比较真实观测与模型投影。",
        ],
    )


def _build_case_slide(slide, config: ReconstructionFlowPptConfig) -> None:
    _add_title(slide, f"当前最小闭环示例：{config.case_id}", "使用现有 clean tree 直接生成 synthetic projection")
    if config.overview_png is not None and config.overview_png.exists():
        slide.shapes.add_picture(str(config.overview_png), Inches(0.7), Inches(1.6), width=Inches(5.5))
    if config.projection_png is not None and config.projection_png.exists():
        slide.shapes.add_picture(str(config.projection_png), Inches(6.7), Inches(1.6), width=Inches(5.1))
    _add_bullet_list(
        slide,
        0.75,
        5.45,
        11.2,
        1.5,
        "当前已实现",
        [
            "左图：stage5 的 clean tree overview；右图：新增的 synthetic X-ray projection preview。",
            "当前支持：世界系放置、简化 C 臂参数、3D 中心线到 2D detector 的投影。",
            "后续可继续接入 ECG、柔性形变、真实 X 光匹配和 Bayesian / MAP 估计。",
        ],
    )


def _build_roadmap_slide(slide, config: ReconstructionFlowPptConfig) -> None:
    _add_title(slide, "可实现路径", "先做最小可行版本，再扩展到术中估计")
    _add_bullet_list(
        slide,
        0.8,
        1.6,
        5.4,
        4.8,
        "近期路线",
        [
            "1. clean tree -> prototype graph：让聚类输出结构先验，而不是只做图像分析。",
            "2. tree -> kinematic tree：引入 q_L / q_R / beta_soft 的参数化表达。",
            "3. synthetic X-ray：采样 C 臂角度、ECG 相位和心脏位姿，生成观测。",
            "4. MAP matching：比较真实 X 光与模型投影，作为 Bayesian filtering 的前一步。",
        ],
    )
    _add_bullet_list(
        slide,
        6.5,
        1.6,
        5.0,
        4.8,
        "建议补充的特征",
        [
            "ancestor / descendant bifurcation count",
            "curvature_max / curvature_p95",
            "5-7 个 shape anchors",
            "role features：trunk / continuation / side / terminal",
        ],
    )
    if config.notes:
        _add_bullet_list(slide, 0.85, 6.0, 10.8, 1.0, "备注", list(config.notes))


def build_reconstruction_flow_ppt(config: ReconstructionFlowPptConfig) -> Path:
    """Generate a compact PPT report for the 3D-2D reconstruction discussion."""
    config.output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _add_title(title_slide, config.title, config.subtitle)
    _add_bullet_list(
        title_slide,
        0.9,
        2.0,
        11.2,
        3.2,
        "核心目标",
        [
            "输入：CCTA / clean tree；输出：可投影、可估计的术中冠脉模型。",
            "思想：把冠脉从静态树升级为参数化状态空间模型。",
            "方法：语义拓扑 + prototype prior + C-arm 投影 + matching / Bayesian 估计。",
        ],
    )

    flow_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _build_flow_slide(flow_slide, config)

    reference_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _build_reference_slide(reference_slide)

    case_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _build_case_slide(case_slide, config)

    roadmap_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    _build_roadmap_slide(roadmap_slide, config)

    presentation.save(str(config.output_path))
    return config.output_path
